import os
import sys
import time
from pathlib import Path
import numpy as np


# ======================================================================
# Utilidades internas
# ======================================================================
def _preguntar_si_no(prompt, default='n'):
    while True:
        try:
            r = input(prompt).strip().lower()
        except EOFError:
            return default == 's'
        if not r:
            return default == 's'
        if r in ('s', 'si', 'sí', 'y', 'yes', '1'):
            return True
        if r in ('n', 'no', '0'):
            return False
        print("  Responde 's' o 'n'.")


def _get_script_info():
    try:
        import __main__
        f = getattr(__main__, "__file__", None)
        if f:
            p = Path(f).resolve()
            return p.parent, p.stem
    except Exception:
        pass
    try:
        p = Path(sys.argv[0]).resolve()
        if p.exists():
            return p.parent, p.stem
    except Exception:
        pass
    return Path.cwd(), "control"


def _sanear_hoja(nombre):
    for ch in ':\\/?*[]':
        nombre = nombre.replace(ch, '-')
    return nombre[:31]


def _icono(ok, warn=False):
    return "✓" if ok else ("~" if warn else "✗")


def _barra(pct, ancho=30):
    llenos = int(round(pct / 100.0 * ancho))
    return "█" * llenos + "·" * (ancho - llenos)


# ======================================================================
# Clase principal
# ======================================================================
class EstadisticasRobustas:
    def __init__(self, fp_target=1.0, tolerancia=0.03, tolerancia_fina=0.01,
                 prefijo_archivo=None):
        self.fp_target = float(fp_target)
        self.tol = float(tolerancia)
        self.tol_fina = float(tolerancia_fina)
        self.prefijo_archivo = prefijo_archivo

        self.t_all, self.fp_all, self.phi_all, self.modo_all = [], [], [], []
        self.t_rango, self.fp_rango, self.phi_rango = [], [], []
        self.fp_rango_fino = []

        self.rachas = []
        self._racha = None
        self.excursiones = []
        self._exc = None

        self.n_total = 0
        self.n_rango = 0
        self.n_rango_fino = 0

        self._cerrado = False
        self._finalizado = False

    # ------------------------------------------------------------------
    # Registro de datos
    # ------------------------------------------------------------------
    def agregar(self, t, fp, phi, modo=''):
        if fp is None or not np.isfinite(fp):
            return
        fp = float(fp)
        phi = float(phi) if (phi is not None and np.isfinite(phi)) else 0.0

        self.n_total += 1
        self.t_all.append(float(t))
        self.fp_all.append(fp)
        self.phi_all.append(phi)
        self.modo_all.append(modo)

        err = abs(fp - self.fp_target)
        en_rango = err <= self.tol

        if en_rango:
            self.n_rango += 1
            self.t_rango.append(float(t))
            self.fp_rango.append(fp)
            self.phi_rango.append(phi)
            if err <= self.tol_fina:
                self.n_rango_fino += 1
                self.fp_rango_fino.append(fp)

            if self._racha is None:
                self._racha = {'t0': float(t), 't_last': float(t), 'n': 1,
                               'fp_min': fp, 'fp_max': fp, 'fp_sum': fp,
                               'phi_max': abs(phi)}
            else:
                r = self._racha
                r['t_last'] = float(t)
                r['n'] += 1
                r['fp_min'] = min(r['fp_min'], fp)
                r['fp_max'] = max(r['fp_max'], fp)
                r['fp_sum'] += fp
                r['phi_max'] = max(r['phi_max'], abs(phi))

            if self._exc is not None:
                self._exc['dur'] = float(t) - self._exc['t0']
                self.excursiones.append(self._exc)
                self._exc = None
        else:
            if self._racha is not None:
                r = self._racha
                r['dur'] = r['t_last'] - r['t0']
                r['fp_mean'] = r['fp_sum'] / max(r['n'], 1)
                self.rachas.append(r)
                self._racha = None

            if self._exc is None:
                self._exc = {'t0': float(t), 'n': 1, 'fp_extremo': fp,
                             'phi_max': abs(phi)}
            else:
                e = self._exc
                e['n'] += 1
                if abs(fp - self.fp_target) > abs(e['fp_extremo'] - self.fp_target):
                    e['fp_extremo'] = fp
                e['phi_max'] = max(e['phi_max'], abs(phi))

    # ------------------------------------------------------------------
    def _cerrar(self, t_final=None):
        if self._cerrado:
            return
        if self._racha is not None:
            r = self._racha
            r['dur'] = r['t_last'] - r['t0']
            r['fp_mean'] = r['fp_sum'] / max(r['n'], 1)
            self.rachas.append(r)
            self._racha = None
        if self._exc is not None:
            self._exc['dur'] = (float(t_final) - self._exc['t0']) \
                if t_final is not None else 0.0
            self.excursiones.append(self._exc)
            self._exc = None
        self._cerrado = True

    # ------------------------------------------------------------------
    def _resumen(self):
        s = {}
        if not self.fp_all:
            return s
        fp_all = np.asarray(self.fp_all)
        err_all = np.abs(fp_all - self.fp_target)

        s['n_total'] = self.n_total
        s['fp_mean_global'] = float(np.mean(fp_all))
        s['fp_std_global'] = float(np.std(fp_all))
        s['fp_min_global'] = float(np.min(fp_all))
        s['fp_max_global'] = float(np.max(fp_all))
        s['fp_median_global'] = float(np.median(fp_all))
        s['err_mean_global'] = float(np.mean(err_all))

        s['n_rango'] = self.n_rango
        s['pct_rango'] = 100.0 * self.n_rango / max(self.n_total, 1)
        s['n_rango_fino'] = self.n_rango_fino
        s['pct_rango_fino'] = 100.0 * self.n_rango_fino / max(self.n_total, 1)

        if self.fp_rango:
            fr = np.asarray(self.fp_rango)
            er = np.abs(fr - self.fp_target)
            s['fp_rango_mean'] = float(np.mean(fr))
            s['fp_rango_std'] = float(np.std(fr))
            s['fp_rango_min'] = float(np.min(fr))
            s['fp_rango_max'] = float(np.max(fr))
            s['fp_rango_span'] = float(np.max(fr) - np.min(fr))
            s['fp_rango_median'] = float(np.median(fr))
            s['fp_rango_p10'] = float(np.percentile(fr, 10))
            s['fp_rango_p90'] = float(np.percentile(fr, 90))
            s['err_rango_max'] = float(np.max(er))
            s['estabilidad_ratio'] = float(np.std(fr) / max(self.tol, 1e-6))

        s['n_rachas'] = len(self.rachas)
        if self.rachas:
            durs = np.asarray([r['dur'] for r in self.rachas])
            ns = np.asarray([r['n'] for r in self.rachas])
            s['racha_dur_max'] = float(np.max(durs))
            s['racha_dur_mean'] = float(np.mean(durs))
            s['racha_dur_median'] = float(np.median(durs))
            s['racha_dur_total'] = float(np.sum(durs))
            s['racha_max_n'] = int(np.max(ns))

            rmax = max(self.rachas, key=lambda r: r['dur'])
            s['racha_max_dur'] = float(rmax['dur'])
            s['racha_max_fp_min'] = float(rmax['fp_min'])
            s['racha_max_fp_max'] = float(rmax['fp_max'])
            s['racha_max_fp_mean'] = float(rmax['fp_mean'])
            s['racha_max_fp_span'] = float(rmax['fp_max'] - rmax['fp_min'])
            s['racha_max_phi_max'] = float(rmax['phi_max'])

        s['n_excursiones'] = len(self.excursiones)
        if self.excursiones:
            edurs = np.asarray([e.get('dur', 0.0) for e in self.excursiones])
            s['exc_dur_max'] = float(np.max(edurs))
            s['exc_dur_mean'] = float(np.mean(edurs))
            s['exc_dur_total'] = float(np.sum(edurs))

        if len(self.fp_all) >= 20:
            m = len(self.fp_all) // 2
            fr2 = np.asarray(self.fp_all[m:])
            er2 = np.abs(fr2 - self.fp_target)
            s['fp_reg_mean'] = float(np.mean(fr2))
            s['fp_reg_std'] = float(np.std(fr2))
            s['fp_reg_min'] = float(np.min(fr2))
            s['fp_reg_max'] = float(np.max(fr2))
            s['pct_reg_rango'] = float(100.0 * np.mean(er2 <= self.tol))

        return s

    # ------------------------------------------------------------------
    # Datos estructurados (para Excel)
    # ------------------------------------------------------------------
    def _filas_estructuradas(self):
        """
        Devuelve una lista de tuplas (Sección, Métrica, Valor, Nota).
        Cada métrica va en su propia celda al exportar.
        """
        s = self._resumen()
        if not s:
            return [("", "Sin datos", "", "")]

        filas = []
        add = lambda sec, met, val, nota='': filas.append((sec, met, val, nota))

        # --- CONFIGURACIÓN ---
        add("CONFIGURACIÓN", "Objetivo FP", self.fp_target)
        add("CONFIGURACIÓN", "Banda objetivo (±)", self.tol)
        add("CONFIGURACIÓN", "Banda fina (±)", self.tol_fina)

        pct_ok, pct_fino = s['pct_rango'], s['pct_rango_fino']
        veredicto = (
            "EXCELENTE" if pct_ok >= 90 else
            "BUENO" if pct_ok >= 75 else
            "ACEPTABLE" if pct_ok >= 50 else
            "NECESITA AJUSTE"
        )

        # --- RESUMEN ---
        add("RESUMEN", "Muestras totales", s['n_total'])
        add("RESUMEN", "Veredicto", veredicto)
        add("RESUMEN", f"Tiempo en ±{self.tol:.2f} (%)", pct_ok)
        add("RESUMEN", f"Tiempo en ±{self.tol_fina:.2f} (%)", pct_fino)

        # --- PRECISIÓN GLOBAL ---
        add("PRECISIÓN GLOBAL", "FP medio", s['fp_mean_global'])
        add("PRECISIÓN GLOBAL", "FP mediana", s['fp_median_global'])
        add("PRECISIÓN GLOBAL", "Desviación std", s['fp_std_global'])
        add("PRECISIÓN GLOBAL", "Error medio", s['err_mean_global'])
        add("PRECISIÓN GLOBAL", "FP mínimo", s['fp_min_global'])
        add("PRECISIÓN GLOBAL", "FP máximo", s['fp_max_global'])

        # --- ESTABILIDAD DENTRO DEL OBJETIVO ---
        if 'fp_rango_mean' in s:
            sec = f"ESTABILIDAD ±{self.tol:.2f}"
            std_v = s['fp_rango_std']
            std_txt = "muy estable" if std_v < 0.005 else \
                      "estable" if std_v < 0.010 else \
                      "algo ruidoso" if std_v < 0.020 else "ruidoso"
            span_v = s['fp_rango_span']
            span_txt = "banda estrecha" if span_v < 0.02 else \
                       "banda media" if span_v < 0.04 else "banda amplia"
            err_max = s['err_rango_max']
            err_txt = "excelente" if err_max < 0.01 else \
                      "bueno" if err_max < 0.02 else "justo"
            ratio = s['estabilidad_ratio']
            ratio_txt = "excelente" if ratio < 0.30 else \
                        "bueno" if ratio < 0.50 else "mejorable"

            add(sec, "Muestras dentro", s['n_rango'])
            add(sec, "Desviación std", std_v, std_txt)
            add(sec, "Amplitud (máx–mín)", span_v, span_txt)
            add(sec, "Error máximo", err_max, err_txt)
            add(sec, "Ratio std/tol", ratio, ratio_txt)
            add(sec, "FP en banda - media", s['fp_rango_mean'])
            add(sec, "FP en banda - P10", s['fp_rango_p10'])
            add(sec, "FP en banda - mediana", s['fp_rango_median'])
            add(sec, "FP en banda - P90", s['fp_rango_p90'])
            add(sec, "FP en banda - mín", s['fp_rango_min'])
            add(sec, "FP en banda - máx", s['fp_rango_max'])

        # --- RACHAS ---
        if s['n_rachas'] == 0:
            add("RACHAS", "Estado", "Nunca dentro de banda")
        else:
            add("RACHAS", "Nº de rachas", s['n_rachas'])
            add("RACHAS", "Duración más larga (s)", s['racha_dur_max'])
            add("RACHAS", "Duración media (s)", s['racha_dur_mean'])
            add("RACHAS", "Duración mediana (s)", s['racha_dur_median'])
            add("RACHAS", "Tiempo total en rachas (s)", s['racha_dur_total'])
            add("RACHAS", "Racha más larga - duración (s)", s['racha_max_dur'])
            add("RACHAS", "Racha más larga - muestras", s['racha_max_n'])
            add("RACHAS", "Racha más larga - FP mín", s['racha_max_fp_min'])
            add("RACHAS", "Racha más larga - FP med", s['racha_max_fp_mean'])
            add("RACHAS", "Racha más larga - FP máx", s['racha_max_fp_max'])
            add("RACHAS", "Racha más larga - amplitud", s['racha_max_fp_span'])
            add("RACHAS", "Racha más larga - |PHI| máx", s['racha_max_phi_max'])

        # --- EXCURSIONES ---
        if s['n_excursiones'] == 0:
            add("EXCURSIONES", "Estado", "Nunca fuera de banda")
        else:
            add("EXCURSIONES", "Nº de salidas", s['n_excursiones'])
            add("EXCURSIONES", "Duración más larga (s)", s['exc_dur_max'])
            add("EXCURSIONES", "Duración media (s)", s['exc_dur_mean'])
            add("EXCURSIONES", "Tiempo total fuera (s)", s['exc_dur_total'])

        # --- RÉGIMEN ---
        if 'fp_reg_mean' in s:
            add("RÉGIMEN", "FP medio", s['fp_reg_mean'])
            add("RÉGIMEN", "Desviación std", s['fp_reg_std'])
            add("RÉGIMEN", "FP mínimo", s['fp_reg_min'])
            add("RÉGIMEN", "FP máximo", s['fp_reg_max'])
            add("RÉGIMEN", "Cobertura en banda (%)", s['pct_reg_rango'])

        # --- VEREDICTO FINAL ---
        if pct_ok >= 90 and s.get('fp_rango_std', 1) < 0.010:
            msg = "CONTROL ESTABLE Y PRECISO — el objetivo se cumple con holgura."
        elif pct_ok >= 75:
            msg = "CONTROL BUENO — mantener dentro de banda con pequeñas mejoras."
        elif pct_ok >= 50:
            msg = "CONTROL ACEPTABLE — revisar ganancias y tiempos de asentamiento."
        else:
            msg = "CONTROL INESTABLE — requiere ajuste de ganancias, slope o deadband."
        add("VEREDICTO FINAL", "Mensaje", msg)

        return filas

    # ------------------------------------------------------------------
    # Reporte de consola (formato bonito)
    # ------------------------------------------------------------------
    def _lineas_reporte(self):
        s = self._resumen()
        if not s:
            return ["[!] Sin datos para el reporte."]

        L = []
        L.append("═" * 78)
        L.append("  REPORTE DE ESTABILIDAD DEL CONTROL")
        L.append(f"  Objetivo: FP = {self.fp_target:.3f}   |   "
                 f"Bandas: ±{self.tol:.2f} (objetivo)  ±{self.tol_fina:.2f} (fino)")
        L.append("═" * 78)
        L.append("")

        pct_ok, pct_fino = s['pct_rango'], s['pct_rango_fino']
        veredicto = (
            f"{_icono(True)}  EXCELENTE" if pct_ok >= 90 else
            f"{_icono(True)}  BUENO"     if pct_ok >= 75 else
            f"{_icono(False, True)}  ACEPTABLE" if pct_ok >= 50 else
            f"{_icono(False)}  NECESITA AJUSTE"
        )
        L.append("▌ RESUMEN")
        L.append(f"  Muestras totales:     {s['n_total']}")
        L.append(f"  Veredicto:            {veredicto}")
        L.append(f"  Tiempo en ±{self.tol:.2f}:     {pct_ok:5.1f}%   {_barra(pct_ok)}")
        L.append(f"  Tiempo en ±{self.tol_fina:.2f}:     {pct_fino:5.1f}%   {_barra(pct_fino)}")
        L.append("")

        L.append("▌ PRECISIÓN GLOBAL  (todas las muestras)")
        L.append(f"  FP medio:            {s['fp_mean_global']:.5f}")
        L.append(f"  FP mediana:          {s['fp_median_global']:.5f}")
        L.append(f"  Desviación std:      {s['fp_std_global']:.5f}")
        L.append(f"  Error medio:         {s['err_mean_global']:.5f}")
        L.append(f"  Rango (min–max):     {s['fp_min_global']:.4f} – {s['fp_max_global']:.4f}")
        L.append("")

        if 'fp_rango_mean' in s:
            std_v = s['fp_rango_std']
            std_txt = "muy estable" if std_v < 0.005 else \
                      "estable" if std_v < 0.010 else \
                      "algo ruidoso" if std_v < 0.020 else "ruidoso"
            span_v = s['fp_rango_span']
            span_txt = "banda estrecha" if span_v < 0.02 else \
                       "banda media" if span_v < 0.04 else "banda amplia"
            err_max = s['err_rango_max']
            err_txt = "excelente" if err_max < 0.01 else \
                      "bueno" if err_max < 0.02 else "justo"
            ratio = s['estabilidad_ratio']
            ratio_txt = "excelente" if ratio < 0.30 else \
                        "bueno" if ratio < 0.50 else "mejorable"

            L.append(f"▌ ESTABILIDAD DENTRO DEL OBJETIVO  (±{self.tol:.2f})")
            L.append(f"  Muestras dentro:  {s['n_rango']}  ({pct_ok:.1f}%)")
            L.append(f"  Desviación std:      {std_v:.5f}   ({std_txt})")
            L.append(f"  Amplitud (máx–mín):  {span_v:.5f}   ({span_txt})")
            L.append(f"  Error máximo:        {err_max:.5f}   ({err_txt})")
            L.append(f"  Ratio std/tol:       {ratio:.3f}     ({ratio_txt})")
            L.append("")
            L.append("  Rango observado en banda:")
            L.append(f"    P10: {s['fp_rango_p10']:.5f}   "
                     f"mediana: {s['fp_rango_median']:.5f}   "
                     f"P90: {s['fp_rango_p90']:.5f}")
            L.append("")

        L.append("▌ RACHAS DENTRO DEL OBJETIVO")
        if s['n_rachas'] == 0:
            L.append(f"  {_icono(False)}  Nunca se mantuvo dentro de la banda.")
        else:
            L.append(f"  Nº de rachas:           {s['n_rachas']}")
            L.append(f"  Duración más larga:     {s['racha_dur_max']:.2f} s")
            L.append(f"  Duración media:         {s['racha_dur_mean']:.2f} s")
            L.append(f"  Duración mediana:       {s['racha_dur_median']:.2f} s")
            L.append(f"  Tiempo total en rachas: {s['racha_dur_total']:.2f} s")
            L.append("")
            L.append(f"  ▸ Racha más larga: {s['racha_max_dur']:.2f} s  "
                     f"({s['racha_max_n']} muestras)")
            L.append(f"      FP min / med / max: "
                     f"{s['racha_max_fp_min']:.5f} / "
                     f"{s['racha_max_fp_mean']:.5f} / "
                     f"{s['racha_max_fp_max']:.5f}")
            L.append(f"      Amplitud: {s['racha_max_fp_span']:.5f}   "
                     f"|PHI| máx: {s['racha_max_phi_max']:.5f}")
        L.append("")

        L.append("▌ EXCURSIONES FUERA DEL OBJETIVO")
        if s['n_excursiones'] == 0:
            L.append(f"  {_icono(True)}  Nunca se salió de la banda objetivo.")
        else:
            L.append(f"  Nº de salidas:          {s['n_excursiones']}")
            L.append(f"  Duración más larga:     {s['exc_dur_max']:.2f} s")
            L.append(f"  Duración media:         {s['exc_dur_mean']:.2f} s")
            L.append(f"  Tiempo total fuera:     {s['exc_dur_total']:.2f} s")
        L.append("")

        if 'fp_reg_mean' in s:
            L.append("▌ RÉGIMEN  (segunda mitad del ensayo)")
            L.append(f"  FP medio:              {s['fp_reg_mean']:.5f}")
            L.append(f"  Desviación std:        {s['fp_reg_std']:.5f}")
            L.append(f"  Rango:                 {s['fp_reg_min']:.5f} – {s['fp_reg_max']:.5f}")
            L.append(f"  Cobertura en banda:    {s['pct_reg_rango']:.1f}%   "
                     f"{_barra(s['pct_reg_rango'])}")
            L.append("")

        if pct_ok >= 90 and s.get('fp_rango_std', 1) < 0.010:
            msg = "CONTROL ESTABLE Y PRECISO — el objetivo se cumple con holgura."
        elif pct_ok >= 75:
            msg = "CONTROL BUENO — mantener dentro de banda con pequeñas mejoras."
        elif pct_ok >= 50:
            msg = "CONTROL ACEPTABLE — revisar ganancias y tiempos de asentamiento."
        else:
            msg = "CONTROL INESTABLE — requiere ajuste de ganancias, slope o deadband."

        L.append("═" * 78)
        L.append(f"  VEREDICTO FINAL: {msg}")
        L.append("═" * 78)
        return L

    # ------------------------------------------------------------------
    # Guardado estructurado en Excel
    # ------------------------------------------------------------------
    def _guardar(self, ruta_excel=None, hoja=None):
        try:
            from openpyxl import Workbook, load_workbook
            from openpyxl.styles import Font, PatternFill, Alignment
            from openpyxl.utils import get_column_letter
        except ImportError:
            print("[X] Falta openpyxl. Instala con: pip install openpyxl")
            return False

        if ruta_excel is None:
            dir_script, script_stem = _get_script_info()
            base = self.prefijo_archivo or script_stem
            ruta_excel = str(Path(dir_script) / f"{base}.xlsx")

        if hoja is None:
            hoja = time.strftime("%Y-%m-%d_%H-%M-%S")
        hoja = _sanear_hoja(hoja)

        if os.path.exists(ruta_excel):
            try:
                wb = load_workbook(ruta_excel)
            except Exception as e:
                print(f"[!] No se pudo abrir {ruta_excel} ({e}). Se crea nuevo.")
                wb = Workbook()
        else:
            wb = Workbook()

        if "Sheet" in wb.sheetnames and len(wb.sheetnames) == 1:
            try:
                del wb["Sheet"]
            except Exception:
                pass

        base_nombre = hoja
        k = 1
        while hoja in wb.sheetnames:
            suf = f"_{k}"
            hoja = base_nombre[:31 - len(suf)] + suf
            k += 1

        ws = wb.create_sheet(title=hoja)

        # --- Encabezados ---
        headers = ["Sección", "Métrica", "Valor", "Nota"]
        f_head = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        fill_head = PatternFill("solid", fgColor="305496")
        align_center = Alignment(horizontal="center", vertical="center")
        for col, h in enumerate(headers, start=1):
            c = ws.cell(row=1, column=col, value=h)
            c.font = f_head
            c.fill = fill_head
            c.alignment = align_center

        # --- Filas (cada métrica en su propia celda) ---
        filas = self._filas_estructuradas()
        f_normal = Font(name="Calibri", size=10)
        f_bold = Font(name="Calibri", size=10, bold=True)
        fill_sec = PatternFill("solid", fgColor="E8EEF5")
        fill_alt = PatternFill("solid", fgColor="F7F9FC")

        seccion_actual = None
        color_toggle = False
        for i, (sec, met, val, nota) in enumerate(filas, start=2):
            # Marcar cambio de sección
            if sec != seccion_actual:
                seccion_actual = sec
                color_toggle = not color_toggle
            fill = fill_alt if color_toggle else None

            c1 = ws.cell(row=i, column=1, value=sec)
            c2 = ws.cell(row=i, column=2, value=met)
            c3 = ws.cell(row=i, column=3, value=val)
            c4 = ws.cell(row=i, column=4, value=nota)

            for c in (c1, c2, c3, c4):
                c.font = f_normal
                if fill is not None:
                    c.fill = fill

            # Sección en negrita en la primera fila de cada bloque
            if i == 2 or (i > 2 and filas[i - 3][0] != sec):
                c1.font = f_bold

            # Formato del valor: si es numérico y pequeño, más decimales
            if isinstance(val, (int, float)) and not isinstance(val, bool):
                if isinstance(val, float) and abs(val) < 1.0:
                    c3.number_format = "0.00000"
                else:
                    c3.number_format = "0.00"

        # Anchos de columnas
        ws.column_dimensions['A'].width = 24
        ws.column_dimensions['B'].width = 34
        ws.column_dimensions['C'].width = 20
        ws.column_dimensions['D'].width = 26
        ws.freeze_panes = "A2"

        try:
            wb.save(ruta_excel)
        except PermissionError:
            print(f"[X] No se puede guardar: {ruta_excel}")
            print("    ¿Está abierto en Excel? Ciérralo e intenta de nuevo.")
            return False
        except Exception as e:
            print(f"[X] Error guardando: {e}")
            return False

        print(f"\n  ✓ Reporte guardado en:")
        print(f"      {ruta_excel}")
        print(f"      Hoja: '{hoja}'")
        return True

    # ------------------------------------------------------------------
    def finalizar(self, t_final=None):
        if self._finalizado:
            return
        self._finalizado = True

        self._cerrar(t_final=t_final)
        for linea in self._lineas_reporte():
            print(linea)

        if not self.fp_all:
            print("[i] No hay muestras para guardar.\n")
            return

        print()
        try:
            quiere = _preguntar_si_no(
                "¿Guardar este reporte en Excel? [s/N]: ", default='n'
            )
        except KeyboardInterrupt:
            print("\n[!] Ctrl+C durante el prompt. Auto-guardando...")
            quiere = True

        if not quiere:
            print("  → No se guardará el reporte.\n")
            return

        self._guardar()


# ======================================================================
# API pública del módulo
# ======================================================================
def solicitar_estadisticas(fp_target=1.0, tolerancia=0.03, tolerancia_fina=0.01,
                           prefijo_archivo=None):
    print()
    usar = _preguntar_si_no(
        "¿Recopilar estadísticas robustas de esta corrida? [s/N]: ",
        default='n'
    )
    if not usar:
        print("  → No se recopilarán estadísticas robustas.\n")
        return None
    print("  → Estadísticas robustas ACTIVADAS.\n")
    return EstadisticasRobustas(
        fp_target=fp_target,
        tolerancia=tolerancia,
        tolerancia_fina=tolerancia_fina,
        prefijo_archivo=prefijo_archivo,
    )


def finalizar_seguro(stats):
    if stats is None:
        return
    try:
        stats.finalizar()
    except KeyboardInterrupt:
        print("\n[!] Ctrl+C durante el cierre. Auto-guardando...")
        try:
            stats._guardar()
        except Exception as e:
            print(f"[X] Falló el auto-guardado: {e}")
    except Exception as e:
        print(f"[!] Error cerrando estadísticas: {e}")
        import traceback
        traceback.print_exc()